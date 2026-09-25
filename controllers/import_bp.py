"""
CONTROLLER (Blueprint) — Import Excel.

Routes:
  GET  /import             -> page
  POST /api/list-sheets    -> list sheet names of uploaded workbook
  POST /import             -> parse + insert into DB
"""
from flask import Blueprint, render_template, request, flash, redirect, url_for, jsonify
import openpyxl

import db
import db_ignore
import db_audit
from services import excel_import

bp = Blueprint('import', __name__)


@bp.route('/import', methods=['GET'])
def import_page():
    return render_template('import.html')


@bp.route('/api/list-sheets', methods=['POST'])
def list_sheets():
    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400
    file = request.files['file']
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'error': 'File harus .xlsx'}), 400
    try:
        wb = openpyxl.load_workbook(file, data_only=True, read_only=True)
        sheets = wb.sheetnames
        wb.close()
        return jsonify({'sheets': sheets})
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@bp.route('/import', methods=['POST'])
def import_excel():
    if 'file' not in request.files:
        flash('Pilih file Excel terlebih dahulu!', 'danger')
        return redirect(url_for('import.import_page'))

    file = request.files['file']
    if not file.filename.endswith(('.xlsx', '.xls')):
        flash('File harus .xlsx atau .xls!', 'danger')
        return redirect(url_for('import.import_page'))

    # ── Project binding (required) ──
    project_id_raw = request.form.get('project_id', '').strip()
    if not project_id_raw:
        flash('Pilih Proyek terlebih dahulu! Import Excel kini terikat ke satu proyek.', 'danger')
        return redirect(url_for('import.import_page'))
    try:
        project_id = int(project_id_raw)
    except (ValueError, TypeError):
        flash('ID Proyek tidak valid.', 'danger')
        return redirect(url_for('import.import_page'))

    # Verify project exists
    import db_project
    project = db_project.get_project(project_id)
    if not project:
        flash('Proyek tidak ditemukan.', 'danger')
        return redirect(url_for('import.import_page'))

    sheet_name = request.form.get('sheet_name', '')
    start_row = int(request.form.get('start_row', 3))

    try:
        wb = openpyxl.load_workbook(file, data_only=True)
    except Exception as e:
        flash(f'Gagal membuka file: {e}', 'danger')
        return redirect(url_for('import.import_page'))

    # Resolve target sheet
    if not sheet_name and 'Monitoring Hutang' in wb.sheetnames:
        sheet_name = 'Monitoring Hutang'
    elif not sheet_name and 'UTANG' in wb.sheetnames:
        sheet_name = 'UTANG'
    elif sheet_name not in wb.sheetnames:
        sheet_name = wb.sheetnames[0]

    ws = wb[sheet_name]

    # ── Preserve ignore & audit references across per-project re-import ──
    # Build (R, rekanan, invoice) → _row map ONLY for rows of THIS project.
    pid_str = str(project_id)
    old_data = db.read_data()
    old_row_to_key = {}
    affected_old_rows = set()
    for r in old_data:
        if str(r.get('project_id', '')) != pid_str:
            continue  # other projects untouched
        k = (
            str(r.get('R', '')),
            (r.get('NAMA LEVELANSIR / REKANAN') or '').strip(),
            (r.get('NO INVOICE') or '').strip(),
        )
        old_row_to_key[r['_row']] = k
        affected_old_rows.add(r['_row'])

    old_ignored_entries = [e for e in db_ignore.get_ignored_rows()
                           if e['hutang_row'] in affected_old_rows]
    old_audit_set = {ar for ar in db_audit.get_audited_rows()
                     if ar in affected_old_rows}

    # Replace only this project's rows; other projects stay intact.
    db.delete_rows_by_project(pid_str)

    result = excel_import.parse_worksheet(ws, start_row, project_id=pid_str)
    if result is None:
        flash('Tidak ada data yang ditemukan untuk diimport!', 'warning')
        wb.close()
        return redirect(url_for('import.import_page'))

    batch, imported, skipped = result
    if batch:
        db.import_rows(batch)
    wb.close()
    from constants import save_last_import as _stamp
    try:
        _stamp('invoice', pid_str)
    except Exception:
        pass

    # ── Relink ignore & audit entries to new _row values (this project only) ──
    new_data = db.read_data()
    new_key_to_row = {}
    for r in new_data:
        if str(r.get('project_id', '')) != pid_str:
            continue
        k = (
            str(r.get('R', '')),
            (r.get('NAMA LEVELANSIR / REKANAN') or '').strip(),
            (r.get('NO INVOICE') or '').strip(),
        )
        new_key_to_row[k] = r['_row']

    row_map = {}
    for entry in old_ignored_entries:
        old_row = entry['hutang_row']
        old_key = old_row_to_key.get(old_row)
        if old_key and old_key in new_key_to_row:
            row_map[old_row] = new_key_to_row[old_key]

    for old_row in old_audit_set:
        if old_row not in row_map:
            old_key = old_row_to_key.get(old_row)
            if old_key and old_key in new_key_to_row:
                row_map[old_row] = new_key_to_row[old_key]

    if row_map:
        db_ignore.relink_ignored(row_map)
        db_audit.relink_audit(row_map)

    flash(f'Import proyek "{project["nama_proyek"]}" selesai! '
          f'{imported} data diimport, {skipped} baris kosong dilewati. '
          f'Proyek lain tidak terpengaruh.', 'success')
    return redirect(url_for('project.project_dashboard', project_id=project_id))
