# -*- coding: utf-8 -*-
"""Smoke check: alur KONFIRMASI import Kontrak (upload -> konfirmasi replace).

Regresi yang dijaga:
  Bug "[Errno 22] Invalid argument" saat klik "Ya, Ganti Data Kontrak".
  Penyebab: stream file unggahan sudah dikonsumsi openpyxl (pointer bergeser),
  lalu file_storage.save() menulis hanya sisa stream -> file temp rusak.

Skrip ini meniru alur nyata:
  1. Upload workbook kontrak MULTI-KB (banyak baris) via POST /kontrak/import
     -> harus membalas halaman konfirmasi (bukan error) dan menyimpan token.
  2. POST konfirmasi (confirmed=1 + import_token) tanpa upload ulang
     -> harus SUKSES (redirect ke /kontrak/), TIDAK [Errno 22].
  3. Kontrak proyek target diganti sesuai isi Excel (jumlah kontrak == Excel).

Isolasi: DB disalin ke folder temp; DATA_DIR diarahkan ke sana sehingga DB
utama TIDAK tersentuh.

Run:  python smoke_contract_import_confirm_check.py
"""
import io
import os
import shutil
import sys
import tempfile

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT)

# ── 1. Siapkan DB temp SEBELUM import app (DATA_DIR dibaca saat import) ──
_TMP = tempfile.mkdtemp(prefix='kontrak_confirm_check_')
_SRC_DB = os.path.join(ROOT, 'monitoring_hutang.db')
if os.path.exists(_SRC_DB):
    shutil.copy2(_SRC_DB, os.path.join(_TMP, 'monitoring_hutang.db'))
os.environ['DATA_DIR'] = _TMP
os.environ.setdefault('DISABLE_EXCEL_AUTOMIGRATE', '1')

import openpyxl  # noqa: E402

# ── 2. Bangun workbook kontrak multi-KB yang valid ──────────────────────
HDR = ['No', 'Tipe Pekerjaan', 'Item Pekerjaan', 'No PO', 'Vendor',
       'Nilai Kontrak', 'Tgl Kontrak Mulai', 'Tgl Selesai', 'PPN (%)',
       'Termin', 'Termin (%)', 'Nilai DPP Termin', 'Nilai PPN Termin',
       'Total DPP', 'Total PPN', 'Metode Pembayaran', 'Durasi Termin',
       'Metode Penetapan', 'Tipe Kontrak', 'Keterangan', 'Lokasi', 'Catatan']


def _build_workbook(n_rows):
    """Buat workbook 'Kontrak Vendor' berisi n_rows baris termin (multi-KB)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Kontrak Vendor'
    for c, h in enumerate(HDR, start=1):
        ws.cell(row=1, column=c, value=h)
    r = 2
    for i in range(n_rows):
        no_po = 'PB.%03d/PO/NK-WIKA/02/SMOKE-CONFIRM/04/2026' % (i + 1)
        nilai = 1_000_000_000 + i * 1_000_000
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value='Pekerjaan Uji')
        ws.cell(row=r, column=3, value='Item uji konfirmasi %d' % (i + 1))
        ws.cell(row=r, column=4, value=no_po)
        ws.cell(row=r, column=5, value='PT. VENDOR SMOKE CONFIRM %d' % (i + 1))
        ws.cell(row=r, column=6, value=nilai)
        ws.cell(row=r, column=7, value='2026-01-05')
        ws.cell(row=r, column=8, value='2026-12-31')
        ws.cell(row=r, column=9, value=11)
        ws.cell(row=r, column=10, value=1)
        ws.cell(row=r, column=11, value=1.0)
        ws.cell(row=r, column=12, value=nilai)
        ws.cell(row=r, column=13, value=0)
        ws.cell(row=r, column=14, value=nilai)
        ws.cell(row=r, column=15, value=0)
        ws.cell(row=r, column=16, value='Transfer')
        ws.cell(row=r, column=17, value=30)
        ws.cell(row=r, column=18, value='Penunjukan Langsung')
        ws.cell(row=r, column=19, value='Lump Sum')
        ws.cell(row=r, column=20, value='Uji regresi konfirmasi')
        ws.cell(row=r, column=21, value='Jakarta')
        ws.cell(row=r, column=22, value='')
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


N_ROWS = 40
xlsx = _build_workbook(N_ROWS)
assert len(xlsx) > 5000, 'workbook uji harus multi-KB (banyak baris) agar bug terpicu'
print('[0] workbook uji: %d baris, %d byte' % (N_ROWS, len(xlsx)))

# ── 3. Jalankan alur via test client ────────────────────────────────────
import app as appmod  # noqa: E402
import db_project  # noqa: E402

app = appmod.app
app.config['WTF_CSRF_ENABLED'] = False
app.config['TESTING'] = True

# Pastikan semua tabel ada (init_db_all hanya jalan di __main__).
appmod.init_db_all()

# Pastikan ada proyek id=1 (dibuat bila DB temp kosong).
if not db_project.get_project(1):
    db_project.create_project('SMOKE1', 'Proyek Smoke Konfirmasi')

client = app.test_client()

# Tandai sudah login (auth = flag session 'authed').
with client.session_transaction() as sess:
    sess['authed'] = True

# --- 3a. Upload (fase 1) harus membalas halaman konfirmasi ---------------
data = {
    'project_id': '1',
    'file': (io.BytesIO(xlsx), 'kontrak_smoke_confirm.xlsx'),
}
resp = client.post('/kontrak/import', data=data,
                   content_type='multipart/form-data')
body = resp.get_data(as_text=True)
assert resp.status_code == 200, 'upload fase-1 status %s' % resp.status_code
assert 'import_token' in body, (
    'halaman konfirmasi tidak muncul (tidak ada import_token). '
    'Mungkin upload gagal/redirect.')
assert 'Ya, Ganti Data Kontrak' in body or 'Ganti Data Kontrak' in body, (
    'tombol konfirmasi tidak ada di halaman')
# Ambil token dari halaman konfirmasi.
import re  # noqa: E402
m = re.search(r'name="import_token"\s+value="([0-9a-f]+)"', body)
assert m, 'token konfirmasi tidak ditemukan di HTML'
token = m.group(1)
print('[1] upload fase-1 OK: halaman konfirmasi tampil, token=%s...' % token[:8])

# --- 3b. Konfirmasi replace (fase 2) harus SUKSES -------------------------
resp2 = client.post('/kontrak/import',
                    data={'confirmed': '1', 'import_token': token})
body2 = resp2.get_data(as_text=True)
# Ikuti redirect agar flash message terbaca.
if resp2.status_code in (301, 302):
    loc = resp2.headers.get('Location', '')
    resp2 = client.get(loc)
    body2 = resp2.get_data(as_text=True)

assert 'Errno 22' not in body2, 'REGRESI: masih [Errno 22] Invalid argument'
assert 'Invalid argument' not in body2, 'REGRESI: file temp rusak saat konfirmasi'
assert 'rusak atau tidak bisa dibaca' not in body2, (
    'REGRESI: file temp tersimpan sebagian (gagal dibaca saat konfirmasi)')
assert 'Import kontrak proyek' in body2, (
    'konfirmasi tidak melaporkan import sukses. Cuplikan: %s'
    % body2[:400].replace('\n', ' '))
print('[2] konfirmasi OK: import sukses, tanpa [Errno 22]')

# --- 3c. Verifikasi hasil: jumlah kontrak proyek 1 == isi Excel ----------
import db_contract  # noqa: E402
rows = db_contract.get_contracts('1')
n_proj = len([c for c in rows if str(c.get('project_id', '')) == '1'])
# Setiap baris termin unik per No PO -> setelah parse jadi N kontrak.
assert n_proj == N_ROWS, (
    f'jumlah kontrak proyek-1 = {n_proj}, diharapkan {N_ROWS} '
    f'(replace per-proyek gagal / parsial)')
print('[3] hasil OK: proyek-1 berisi %d kontrak (sesuai Excel)' % n_proj)

print('ALL CHECKS PASSED')
shutil.rmtree(_TMP, ignore_errors=True)
