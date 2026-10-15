# -*- coding: utf-8 -*-
"""Smoke check: import Kontrak Excel (parser services/contract_excel.py).

Fokus: file procurement nyata "docs/Rekap Kontrak Vendor Procurement Only.xlsx"
memakai header "Tgl Kontrak Mulai" (bukan "Tgl Kontrak"). Regresi pernah
membuat kolom tanggal kontrak tidak terbaca (tgl_kontrak kosong). Skrip ini
memastikan:
  1. File nyata diparse: tgl_kontrak terisi untuk SEMUA kontrak.
  2. Format lama (docs/Template Kontrak Vendor.xlsx) tetap dikenali.
  3. Template yang di-download strukturnya identik dgn file nyata dan
     header-nya bisa di-round-trip (generate -> parse) dengan tanggal utuh.

Run:  python smoke_contract_import_check.py
Read-only: hanya membaca file di docs/ dan workbook in-memory.
"""
import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl

from services import contract_excel as ce

ROOT = os.path.dirname(os.path.abspath(__file__))


def _parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[ce.detect_contract_sheet(wb)]
    res = ce.parse_contract_sheet(ws, project_id='1')
    wb.close()
    return res


# --- 1. File procurement nyata: tgl_kontrak WAJIB terisi semua -------------
real = os.path.join(ROOT, 'docs', 'Rekap Kontrak Vendor Procurement Only.xlsx')
assert os.path.exists(real), f'file tidak ditemukan: {real}'
rows, imported, skipped = _parse(real)
assert rows, 'tidak ada kontrak terbaca dari file nyata'
n_tgl = sum(1 for r in rows if r['tgl_kontrak'])
n_selesai = sum(1 for r in rows if r['tgl_selesai'])
assert n_tgl == len(rows), (
    f'regresi: tgl_kontrak terisi {n_tgl}/{len(rows)} — header '
    f'"Tgl Kontrak Mulai" tidak dikenali parser')
assert n_selesai == len(rows), f'tgl_selesai terisi {n_selesai}/{len(rows)}'
assert all(r['vendor'] for r in rows), 'ada kontrak tanpa vendor di file nyata'
print(f'[1] file nyata OK: {len(rows)} kontrak, {imported} baris, '
      f'tgl_kontrak {n_tgl}/{len(rows)}, tgl_selesai {n_selesai}/{len(rows)}')

# --- 2. Format lama tetap dikenali (tidak regresi) -------------------------
legacy = os.path.join(ROOT, 'docs', 'Template Kontrak Vendor.xlsx')
if os.path.exists(legacy):
    rows_l, imported_l, _ = _parse(legacy)
    assert rows_l, 'format lama tidak terbaca'
    n_tgl_l = sum(1 for r in rows_l if r['tgl_kontrak'])
    assert n_tgl_l == len(rows_l), (
        f'regresi format lama: tgl_kontrak {n_tgl_l}/{len(rows_l)}')
    print(f'[2] format lama OK: {len(rows_l)} kontrak, '
          f'tgl_kontrak {n_tgl_l}/{len(rows_l)}')
else:
    print('[2] file format lama tidak ada, dilewati')

# --- 3. Template download identik dgn file nyata + round-trip --------------
import app as appmod  # noqa: E402  (butuh sys.path di atas)

app = appmod.app
with app.test_request_context('/'):
    from controllers.contract_bp import import_template
    resp = import_template()
    resp.direct_passthrough = False
    data = resp.get_data()

wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
assert 'Kontrak Vendor' in wb.sheetnames, 'sheet "Kontrak Vendor" hilang'
ws = wb['Kontrak Vendor']

# Header template harus sama persis dgn file nyata (22 kolom fungsional).
wb_real = openpyxl.load_workbook(real, data_only=True)
ws_real = wb_real['Kontrak PO Vendor']
hdr_tpl = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
hdr_real = [ws_real.cell(row=1, column=c).value
            for c in range(1, 23)]  # file nyata: 22 kolom fungsional pertama
assert hdr_tpl == hdr_real, (
    'header template download tidak sama dengan file procurement nyata:\n'
    f'  template: {hdr_tpl}\n  nyata   : {hdr_real}')

# Semua field wajib parser harus dikenali dari header template.
_, cmap = ce._find_header(ws)
required = {'no_kontrak', 'vendor', 'nilai', 'tgl_kontrak', 'tgl_selesai',
            'termin', 'termin_persen', 'termin_nilai', 'ppn_nilai',
            'total_dpp', 'total_ppn', 'ppn_rate', 'metode_pembayaran',
            'durasi_termin', 'metode_penetapan', 'tipe_kontrak'}
missing = required - set(cmap)
assert not missing, f'header template tidak dikenali parser: {missing}'

# Round-trip: salin baris contoh sheet Petunjuk -> sheet Kontrak Vendor,
# pastikan tanggal kontrak tetap terbaca.
guide = wb['Petunjuk']
hdr_row = next(rr for rr in range(1, guide.max_row + 1)
               if guide.cell(row=rr, column=4).value == 'No PO')
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = 'Kontrak Vendor'
for c in range(1, 23):
    ws2.cell(row=1, column=c, value=guide.cell(row=hdr_row, column=c).value)
r2 = 2
for rr in range(hdr_row + 1, guide.max_row + 1):
    vals = [guide.cell(row=rr, column=c).value for c in range(1, 23)]
    if all(v in (None, '') for v in vals):
        continue
    for c, v in enumerate(vals, 1):
        ws2.cell(row=r2, column=c, value=v)
    r2 += 1
rows_rt, _, _ = ce.parse_contract_sheet(ws2, project_id='1')
assert rows_rt, 'round-trip template tidak menghasilkan kontrak'
assert all(r['tgl_kontrak'] for r in rows_rt), (
    'round-trip: tgl_kontrak kosong — template contoh tidak konsisten')
print(f'[3] template download OK: {len(hdr_tpl)} kolom identik, '
      f'round-trip {len(rows_rt)} kontrak tanggal utuh')

print('ALL CHECKS PASSED')
